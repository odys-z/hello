"""Builds samples/tricky.docx: headings detectable only via style inheritance,
direct outlineLvl, a content control, a numbered list and a merged-cell table."""
import sys
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

out = sys.argv[1]
d = Document()
st = d.styles

# custom style whose name says nothing, but basedOn Heading 1 -> inherits outline level
ch = st.add_style('章标题', WD_STYLE_TYPE.PARAGRAPH); ch.base_style = st['Heading 1']
# custom style with its own outlineLvl=2 (would be level 3), not named "heading"
cl = st.add_style('Clause', WD_STYLE_TYPE.PARAGRAPH)
cl.element.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="2"/>'))

d.add_paragraph('Annual Report', style='Title')
d.add_paragraph('第一章 总则', style='章标题')
d.add_paragraph('Body text under chapter one.')
d.add_paragraph('1.1 Scope', style='Heading 2')
p = d.add_paragraph('Direct outline level on a Normal paragraph')
p._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="1"/>'))
d.add_paragraph('Clause 7: confidentiality', style='Clause')
d.add_paragraph('First bullet', style='List Number')
d.add_paragraph('Second bullet', style='List Number')

t = d.add_table(rows=3, cols=3); t.style = 'Table Grid'
for r in range(3):
    for c in range(3): t.cell(r, c).text = f'r{r}c{c}'
t.cell(0, 0).merge(t.cell(0, 1)).text = 'merged header'
t.cell(1, 2).merge(t.cell(2, 2)).text = 'vertical merge'

# heading wrapped in a content control (w:sdt) - many readers skip these
sdt = parse_xml(f'''<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="Section"/></w:sdtPr>
  <w:sdtContent><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
  <w:r><w:t>第二章 Inside a content control</w:t></w:r></w:p></w:sdtContent></w:sdt>''')
d.element.body.insert(len(d.element.body) - 1, sdt)
d.add_paragraph('Body text under chapter two.')
d.save(out)
